"""Importa categorias da planilha para meta/mods.toml (spec §2.5, §3).

Uso: importar_planilha.py XLSX [--pack DIR]
Lê a aba "Testes de Mods" (colunas `Mod`, `Categoria`, `Cliente ou Servidor?`), casa cada linha
com um `.pw.toml` (ou item de manual.toml) por nome normalizado e escreve/mescla
`meta/mods.toml` sem apagar `classe`/`licenca`/`nota` já editados. O que não casar vai para
`meta/sem-par.txt` (apagado se vazio).
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import openpyxl
import tomli_w

import _comum as c

ABA = "Testes de Mods"
COL_MOD, COL_CAT, COL_LADO = "Mod", "Categoria", "Cliente ou Servidor?"
PRESERVAR = ("classe", "licenca", "nota")

# §2.5 — chave = texto normalizado da planilha
MAPA_CATEGORIA = {
    c.normalizar(k): v
    for k, v in {
        "Anime / Chefes / Magia": "tensura",
        "Magia / Combate / Chefes": "magia",
        "Tecnologia": "tecnologia",
        "NPCs / Construção": "colonias",
        "Exploração / Biomas": "mundo",
        "Construção / Decoração": "construcao",
        "Imersão / Visual": "sobrevivencia",
        "Qualidade de Vida": "qol",
        "Desempenho": "performance",
        "Bibliotecas": "bibliotecas",
    }.items()
}


def categoria_slug(texto: str | None) -> str | None:
    """Slug da categoria; None se desconhecida/vazia (o chamador usa `outros` + aviso)."""
    return MAPA_CATEGORIA.get(c.normalizar(texto or "")) if texto else None


def ler_planilha(xlsx: Path) -> list[dict]:
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    if ABA not in wb.sheetnames:
        raise ValueError(f"aba \"{ABA}\" não encontrada em {xlsx}")
    linhas = list(wb[ABA].iter_rows(values_only=True))
    wb.close()
    for i, cab in enumerate(linhas):
        nomes = [str(x).strip() if x is not None else "" for x in cab]
        if COL_MOD in nomes and COL_CAT in nomes:
            break
    else:
        raise ValueError(f"cabeçalho com \"{COL_MOD}\" e \"{COL_CAT}\" não encontrado")
    idx = {n: j for j, n in enumerate(nomes)}
    saida = []
    for r in linhas[i + 1:]:
        def col(nome: str) -> str:
            j = idx.get(nome)
            v = r[j] if j is not None and j < len(r) else None
            return "" if v is None else str(v).strip()

        if col(COL_MOD):
            saida.append({"mod": col(COL_MOD), "categoria": col(COL_CAT), "lado": col(COL_LADO)})
    return saida


def importar(xlsx: Path, pack_dir: Path) -> tuple[dict[str, dict], list[str]]:
    """→ (meta mesclado, linhas sem par). Grava meta/mods.toml e meta/sem-par.txt."""
    pack_dir = Path(pack_dir)
    mods = c.ler_mods(pack_dir)
    manual = c.ler_manual(pack_dir)
    meta = {k: dict(v) for k, v in c.ler_meta(pack_dir).items()}

    candidatos: dict[str, tuple[str, str, str]] = {}  # nome normalizado → (id, nome, classe)
    for m in mods:
        classe = "modrinth" if m.fonte == "modrinth" else ""
        for chave in (c.normalizar(m.id), c.normalizar(m.nome)):
            candidatos.setdefault(chave, (m.id, m.nome, classe))
    for item in manual:
        mid = str(item.get("id", ""))
        nome = str(item.get("nome", mid))
        for chave in (c.normalizar(mid), c.normalizar(nome)):
            candidatos.setdefault(chave, (mid, nome, str(item.get("classe", ""))))

    sem_par: list[str] = []
    for linha in ler_planilha(xlsx):
        alvo = candidatos.get(c.normalizar(linha["mod"]))
        if alvo is None:
            sem_par.append(f"{linha['mod']}\t{linha['categoria']}\t{linha['lado']}")
            continue
        mod_id, nome, classe_padrao = alvo
        cat = categoria_slug(linha["categoria"])
        if cat is None:
            print(
                f"aviso: {mod_id}: categoria \"{linha['categoria']}\" desconhecida → outros",
                file=sys.stderr,
            )
            cat = "outros"
        entrada = meta.setdefault(mod_id, {})
        entrada.setdefault("nome", nome)
        entrada["categoria"] = cat
        entrada["planilha"] = linha["categoria"]
        if classe_padrao:
            entrada.setdefault("classe", classe_padrao)
        # classe/licenca/nota já presentes nunca são sobrescritos (setdefault acima)

    destino = pack_dir / "meta" / "mods.toml"
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_text(tomli_w.dumps(dict(sorted(meta.items()))), encoding="utf-8")
    arq_sem_par = pack_dir / "meta" / "sem-par.txt"
    if sem_par:
        arq_sem_par.write_text("\n".join(sem_par) + "\n", encoding="utf-8")
    elif arq_sem_par.exists():
        arq_sem_par.unlink()
    return meta, sem_par


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Importa a planilha para meta/mods.toml.")
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("--pack", type=Path, default=c.PACK_DIR)
    args = ap.parse_args(argv)
    try:
        meta, sem_par = importar(args.xlsx, args.pack)
    except (ValueError, FileNotFoundError) as e:
        print(f"erro: {e}", file=sys.stderr)
        return 1
    print(f"{len(meta)} mod(s) em meta/mods.toml; {len(sem_par)} sem par (meta/sem-par.txt).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
