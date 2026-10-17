"""CA-T9 — importar_planilha.py (planilha de 5 linhas criada pelo teste com openpyxl)."""

from __future__ import annotations

import tomllib

import openpyxl
import pytest

import importar_planilha as imp

LINHAS = [
    ("Create", "Tecnologia", "Ambos (obrigatório)"),
    ("Just Enough Items (JEI)", "Qualidade de Vida", "Qualquer lado"),
    ("MineColonies", "NPCs / Construção", "Ambos"),
    ("Xaero's World Map", "Coisa Estranha", "Cliente"),
    ("Mod Que Não Existe", "Magia / Combate / Chefes", "Ambos"),
]


@pytest.fixture
def xlsx(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Testes de Mods"
    ws.append(["Mod", "Categoria", "Compatível?", "Cliente ou Servidor?"])
    for mod, cat, lado in LINHAS:
        ws.append([mod, cat, "Pendente", lado])
    p = tmp_path / "planilha.xlsx"
    wb.save(p)
    return p


def _meta(pack):
    return tomllib.loads((pack / "meta/mods.toml").read_text())


@pytest.fixture
def pack_sem_meta(pack_ok):
    (pack_ok / "meta/mods.toml").unlink()
    return pack_ok


def test_mapa_e_sem_par(pack_sem_meta, xlsx, capsys):
    assert imp.main([str(xlsx), "--pack", str(pack_sem_meta)]) == 0
    meta = _meta(pack_sem_meta)
    assert meta["create"] == {"nome": "Create", "categoria": "tecnologia",
                              "planilha": "Tecnologia", "classe": "modrinth"}
    assert meta["minecolonies"]["categoria"] == "colonias"
    assert "classe" not in meta["minecolonies"]  # CF: classe é decisão manual (A/B)
    # categoria desconhecida → outros + aviso
    assert meta["xaeros-world-map"]["categoria"] == "outros"
    assert meta["xaeros-world-map"]["planilha"] == "Coisa Estranha"
    assert "Coisa Estranha" in capsys.readouterr().err
    # "Just Enough Items (JEI)" não casa com "Just Enough Items" → sem par
    sem_par = (pack_sem_meta / "meta/sem-par.txt").read_text().splitlines()
    assert [s.split("\t")[0] for s in sem_par] == ["Just Enough Items (JEI)", "Mod Que Não Existe"]
    assert "jei" not in meta


def test_todas_as_categorias_do_mapa():
    esperado = {
        "Anime / Chefes / Magia": "tensura", "Magia / Combate / Chefes": "magia",
        "Tecnologia": "tecnologia", "NPCs / Construção": "colonias",
        "Exploração / Biomas": "mundo", "Construção / Decoração": "construcao",
        "Imersão / Visual": "sobrevivencia", "Qualidade de Vida": "qol",
        "Desempenho": "performance", "Bibliotecas": "bibliotecas",
    }
    for texto, s in esperado.items():
        assert imp.categoria_slug(texto) == s
    assert imp.categoria_slug("") is None and imp.categoria_slug("Outra") is None


def test_segunda_rodada_preserva_edicoes(pack_ok, xlsx):
    t = (pack_ok / "meta/mods.toml").read_text()
    t = t.replace('[minecolonies]\nnome = "MineColonies"\ncategoria = "colonias"',
                  '[minecolonies]\nnome = "MineColonies"\ncategoria = "mundo"\nlicenca = "MIT"')
    t = t.replace('classe = "modrinth"\nnota = ""\n\n[ftb-library]',
                  'classe = "modrinth"\nnota = "obrigatório no cliente"\n\n[ftb-library]')
    (pack_ok / "meta/mods.toml").write_text(t)

    imp.importar(xlsx, pack_ok)
    imp.importar(xlsx, pack_ok)  # idempotente
    meta = _meta(pack_ok)
    assert meta["minecolonies"]["classe"] == "A"
    assert meta["minecolonies"]["licenca"] == "MIT"
    assert meta["minecolonies"]["categoria"] == "colonias"  # categoria vem da planilha
    assert meta["create"]["nota"] == "obrigatório no cliente"
    assert meta["ftb-library"]["classe"] == "B"  # fora da planilha: intocado
    assert set(meta) >= {"ftb-teams", "slime-gui-pack", "jei"}


def test_sem_par_vazio_apaga_arquivo(pack_ok, tmp_path):
    (pack_ok / "meta/sem-par.txt").write_text("velho\n")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Testes de Mods"
    ws.append(["Mod", "Categoria", "Cliente ou Servidor?"])
    ws.append(["FTB Teams", "Bibliotecas", "Ambos"])  # casa com item C do manual.toml
    p = tmp_path / "um.xlsx"
    wb.save(p)
    _, sem_par = imp.importar(p, pack_ok)
    assert sem_par == [] and not (pack_ok / "meta/sem-par.txt").exists()


def test_aba_errada_sai_1(pack_ok, tmp_path, capsys):
    wb = openpyxl.Workbook()
    wb.active.title = "Outra"
    p = tmp_path / "x.xlsx"
    wb.save(p)
    assert imp.main([str(p), "--pack", str(pack_ok)]) == 1
    assert "Testes de Mods" in capsys.readouterr().err
